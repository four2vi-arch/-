#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""돌다리 v1.0 · 가상 기관 모의 업무폴더와 정답표 만들기

가상 기관 「한빛군청 행정지원과 서무담당」의 1년치(2025.3.~2026.2.) 업무 자료를 합성한다.
실제 기관·사람·전화·문서는 하나도 쓰지 않는다(전화는 010-0000-/063-000-, 메일은 example.kr).
만드는 것: 한글(.hwpx) 공문 12 · 엑셀 5 · 워드 4 · PDF 3 · 메일 4 · 메모 4 · 중첩 압축 1 · 못 읽을 파일 3 = 36개
정답표(정답표.xlsx)는 같은 자료에서 자동으로 만들므로 발췌가 원문과 글자 단위로 일치한다.

쓰임: python make_mock.py <출력 폴더> [hwpx 틀 폴더]
  hwpx 틀 폴더: 실제 한글이 만든 hwpx를 푼 폴더(Contents/header.xml 등). 없으면 .hwpx 대신 .txt로 만든다.
PDF는 LibreOffice(soffice)가 있으면 워드에서 변환하고, 없으면 건너뛴다.
"""
import os, sys, re, io, json, csv, zipfile, random, shutil, subprocess, email.message, email.utils, datetime
from xml.sax.saxutils import escape

ORG = '한빛군청'
DEPT = '행정지원과'
ME = '서무담당(전임자 가상 이름: 오지영)'
random.seed(20261009)

# ---------------------------------------------------------------- 가상 협의 상대(모두 가상)
PEOPLE = [
    ('강민재', '한빛군청 재정과 예산담당', '063-000-0101', '예산 편성·추경 협의'),
    ('이수현', '한빛군청 총무과 청사관리담당', '063-000-0202', '청사 시설·소방 점검 일정 협의'),
    ('박도윤', '한빛소방서 예방안전과', '063-000-0303', '소방안전점검 입회'),
    ('정하은', '한빛군의회 사무과', '063-000-0404', '행정사무감사 자료 요구'),
    ('최우진', '주식회사 가온복합기(가상 업체) 영업팀', '010-0000-0505', '복합기 임대 계약·소모품'),
    ('한서윤', '한빛군청 정보통신과 보안담당', '063-000-0606', '개인정보 보호 자체점검'),
    ('문지호', '전북특별자치도 자치행정과(가상 부서명)', '063-000-0707', '기록물 이관 지침 문의'),
    ('윤채원', '한빛군청 감사실', '063-000-0808', '자체감사 지적사항 조치'),
]

# ---------------------------------------------------------------- 월별 업무(공문 12건의 뼈대)
# (월, 문서번호, 제목, 본문 문장들, 기한, 협의자 색인, 할 일 요약)
MONTHS = [
    (3, '행정지원과-1023', '2025년 상반기 소모품 구매 계획 수립',
     ['부서 소모품 수요를 3월 20일까지 취합하여 상반기 구매 계획을 수립합니다.', '구매는 조달청 나라장터 종합쇼핑몰을 이용하며 1회 구매 한도는 300만 원입니다.', '수요조사 결과는 소모품 구매대장에 기록합니다.'],
     '2025-03-20', 0, '3월 20일까지 부서 소모품 수요 취합 후 상반기 구매 계획 수립'),
    (4, '행정지원과-1388', '2025년 청사 소방안전점검 실시 계획',
     ['4월 셋째 주에 한빛소방서 입회 아래 청사 소방안전점검을 실시합니다.', '점검 결과 지적사항은 4월 30일까지 조치하고 총무과에 통보합니다.', '점검 일정은 소방서 예방안전과와 전화로 조율합니다.'],
     '2025-04-30', 2, '4월 셋째 주 소방안전점검 실시, 지적사항 4월 30일까지 조치'),
    (5, '행정지원과-1720', '복합기 임대계약 갱신 검토',
     ['복합기 임대계약(계약번호 HB-2023-07)이 2025년 6월 30일 만료됩니다.', '5월 말까지 사용량과 장애 이력을 정리하여 갱신 여부를 결정합니다.', '갱신 시 임대료는 월 180,000원을 넘지 않도록 협의합니다.'],
     '2025-05-31', 4, '5월 말까지 복합기 임대계약 갱신 여부 결정(6월 30일 만료)'),
    (6, '행정지원과-2015', '2025년 개인정보 보호 자체점검 실시',
     ['6월 10일부터 6월 24일까지 부서 개인정보 처리 현황을 자체점검합니다.', '점검표는 정보통신과 보안담당이 배포하며 결과는 6월 27일까지 제출합니다.', '개인정보 파일 대장을 최신으로 갱신합니다.'],
     '2025-06-27', 5, '6월 10일~24일 개인정보 자체점검, 결과 6월 27일까지 제출'),
    (7, '행정지원과-2301', '2026년도 본예산 편성 요구서 제출',
     ['2026년도 본예산 편성 요구서를 7월 18일까지 재정과에 제출합니다.', '사무관리비와 공공운영비는 전년 대비 3% 이내 증액을 원칙으로 합니다.', '증액 사유가 있는 사업은 별도 설명 자료를 첨부합니다.'],
     '2025-07-18', 0, '7월 18일까지 2026년도 본예산 편성 요구서 제출'),
    (8, '행정지원과-2654', '여름철 청사 냉방 운영 및 전기요금 절감 안내',
     ['8월 한 달간 청사 냉방 설정 온도를 26도로 유지합니다.', '점심시간 소등과 퇴근 시 사무기기 전원 차단을 실시합니다.', '월별 전기 사용량은 총무과 청사관리담당이 집계합니다.'],
     '2025-08-31', 1, '8월 냉방 26도 유지, 전기 사용량 총무과와 집계'),
    (9, '행정지원과-2980', '2025년 행정사무감사 자료 제출',
     ['군의회 행정사무감사 요구자료를 9월 12일까지 의회 사무과에 제출합니다.', '자료 목록은 사무관리비 집행내역, 계약 현황, 소모품 구매대장입니다.', '제출 전 과장 검토를 거칩니다.'],
     '2025-09-12', 3, '9월 12일까지 행정사무감사 요구자료 제출'),
    (10, '행정지원과-3217', '2025년 자체감사 지적사항 조치 계획',
     ['자체감사에서 지적된 소모품 구매대장 기재 누락 3건을 10월 15일까지 보완합니다.', '조치 결과는 감사실에 서면으로 통보합니다.', '재발 방지를 위해 월 1회 대장 점검을 정례화합니다.'],
     '2025-10-15', 7, '10월 15일까지 자체감사 지적사항(대장 누락 3건) 보완'),
    (11, '행정지원과-3555', '2025년 기록물 정리 및 이관 계획',
     ['보존기간이 지난 기록물을 11월 28일까지 정리하여 기록관에 이관합니다.', '이관 목록은 기록물 관리 지침에 따라 작성합니다.', '지침 해석이 필요한 사항은 도 자치행정과에 문의합니다.'],
     '2025-11-28', 6, '11월 28일까지 기록물 정리·기록관 이관'),
    (12, '행정지원과-3890', '연말 재물조사 실시 계획',
     ['12월 8일부터 12월 12일까지 부서 비품 재물조사를 실시합니다.', '조사 결과는 12월 19일까지 재정과에 제출합니다.', '불용 비품은 처분 목록을 따로 만듭니다.'],
     '2025-12-19', 0, '12월 8일~12일 재물조사, 결과 12월 19일까지 제출'),
    (1, '행정지원과-0112', '2026년 상반기 청사 소방안전점검 사전 협의',
     ['2026년 상반기 소방안전점검은 4월 중 실시 예정이며 1월 말까지 소방서와 일정을 사전 협의합니다.', '지난해 지적사항 조치 결과를 함께 보고합니다.', '점검 전 소화기 충약 상태를 확인합니다.'],
     '2026-01-30', 2, '1월 말까지 상반기 소방안전점검 일정 사전 협의'),
    (2, '행정지원과-0455', '2026년 사무관리비 집행 계획 수립',
     ['2026년 사무관리비 집행 계획을 2월 20일까지 수립합니다.', '분기별 집행 목표는 25%로 하고 4분기 쏠림을 막습니다.', '집행 현황은 매월 말 예산집행현황 파일에 갱신합니다.'],
     '2026-02-20', 0, '2월 20일까지 2026년 사무관리비 집행 계획 수립'),
]

# ---------------------------------------------------------------- 진행 중 현안(메모·메일·워드에 흩어짐)
ISSUES = [
    ('복합기 임대계약 갱신 미결', '업체가 월 임대료 195,000원을 요구하여 협의가 끝나지 않았습니다.', '메모'),
    ('기록물 이관 목록 반려', '기록관에서 이관 목록 서식 오류로 반려되어 재작성이 필요합니다.', '메일'),
    ('자체감사 지적사항 1건 미조치', '소모품 구매대장 누락 3건 중 1건은 영수증을 찾지 못해 미조치 상태입니다.', '워드'),
    ('예산 집행률 저조', '2025년 12월 기준 사무관리비 집행률이 71%로 목표에 미달합니다.', '엑셀'),
    ('청사 2층 복도 누수', '2026년 1월부터 2층 복도 천장 누수가 반복되어 총무과에 보수를 요청해 두었습니다.', '메모'),
]

ANSWERS = []  # (유형, 내용, 시기, 출처 파일, 출처 위치, 발췌)


def add_answer(kind, text, when, file, loc, quote):
    ANSWERS.append([kind, text, when, file, loc, quote])


# ---------------------------------------------------------------- 파일 만들기
def write_hwpx(tpl, path, paras):
    if not tpl:
        open(path.replace('.hwpx', '.txt'), 'w', encoding='utf-8').write('\n'.join(paras))
        return os.path.basename(path).replace('.hwpx', '.txt')
    sec = open(os.path.join(tpl, 'Contents', 'section0.xml'), encoding='utf-8').read()
    head = sec[:sec.index('<hp:p ')]
    m = re.search(r'<hp:p [^>]*>.*?</hp:p>', sec, flags=re.S)
    first = re.sub(r'<hp:t>.*?</hp:t>', '', m.group(0), flags=re.S)
    body = ''.join('<hp:p id="%d" paraPrIDRef="84" styleIDRef="0" pageBreak="0" columnBreak="0" merged="0"><hp:run charPrIDRef="105"><hp:t>%s</hp:t></hp:run></hp:p>' % (i + 1, escape(t)) for i, t in enumerate(paras))
    xml = head + first + body + '</hs:sec>'
    with zipfile.ZipFile(path, 'w') as z:
        z.writestr(zipfile.ZipInfo('mimetype'), 'application/hwp+zip', compress_type=zipfile.ZIP_STORED)
        for root, dirs, files in os.walk(tpl):
            for f in sorted(files):
                p = os.path.join(root, f)
                rel = os.path.relpath(p, tpl).replace(os.sep, '/')
                if rel == 'mimetype':
                    continue
                if rel == 'Contents/section0.xml':
                    z.writestr(rel, xml.encode('utf-8'), compress_type=zipfile.ZIP_DEFLATED)
                elif rel == 'Preview/PrvText.txt':
                    z.writestr(rel, '\n'.join(paras).encode('utf-16'), compress_type=zipfile.ZIP_DEFLATED)
                elif rel.startswith('Preview/'):
                    continue
                else:
                    z.write(p, rel, compress_type=zipfile.ZIP_DEFLATED)
    return os.path.basename(path)


def write_docx(path, paras, table=None):
    W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    ps = ''.join('<w:p><w:r><w:t xml:space="preserve">%s</w:t></w:r></w:p>' % escape(t) for t in paras)
    tbl = ''
    if table:
        rows = ''.join('<w:tr>' + ''.join('<w:tc><w:tcPr><w:tcW w:w="2400" w:type="dxa"/></w:tcPr><w:p><w:r><w:t xml:space="preserve">%s</w:t></w:r></w:p></w:tc>' % escape(c) for c in row) + '</w:tr>' for row in table)
        tbl = '<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/><w:tblW w:w="0" w:type="auto"/><w:tblBorders><w:top w:val="single" w:sz="4"/><w:left w:val="single" w:sz="4"/><w:bottom w:val="single" w:sz="4"/><w:right w:val="single" w:sz="4"/><w:insideH w:val="single" w:sz="4"/><w:insideV w:val="single" w:sz="4"/></w:tblBorders></w:tblPr>' + rows + '</w:tbl><w:p/>'
    doc = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="%s"><w:body>%s%s<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="708" w:footer="708" w:gutter="0"/></w:sectPr></w:body></w:document>' % (W, ps, tbl)
    ct = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
    rels = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', ct)
        z.writestr('_rels/.rels', rels)
        z.writestr('word/document.xml', doc)


def write_xlsx(path, sheets):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(path)


def write_eml(path, frm, to, date, subject, body, attach=None):
    msg = email.message.EmailMessage()
    msg['From'] = frm
    msg['To'] = to
    msg['Date'] = email.utils.format_datetime(date)
    msg['Subject'] = subject
    msg.set_content(body)
    if attach:
        msg.add_attachment(attach[1], maintype='application', subtype='octet-stream', filename=attach[0])
    with open(path, 'wb') as f:
        f.write(msg.as_bytes())


def docx_to_pdf(docx_path, out_dir):
    if not shutil.which('soffice'):
        return None
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', out_dir, docx_path], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    pdf = os.path.join(out_dir, os.path.splitext(os.path.basename(docx_path))[0] + '.pdf')
    return pdf if os.path.exists(pdf) else None


# ---------------------------------------------------------------- 본체
def main():
    out = sys.argv[1]
    tpl = sys.argv[2] if len(sys.argv) > 2 and os.path.isdir(os.path.join(sys.argv[2], 'Contents')) else None
    if os.path.exists(out):
        shutil.rmtree(out)
    folder = os.path.join(out, '가상기관_모의업무폴더')
    for d in ['01_공문', '02_대장·현황(엑셀)', '03_계획·매뉴얼(워드)', '04_메일', '05_메모', '06_기타']:
        os.makedirs(os.path.join(folder, d))
    made = []

    # 1) 공문 hwpx 12
    for (mon, no, title, sents, due, pi, todo) in MONTHS:
        year = 2025 if mon >= 3 else 2026
        date = '%d. %d. %d.' % (year, mon, random.randint(2, 9))
        who = PEOPLE[pi]
        paras = [ORG, '수신 내부결재', '(경유)', '제목 ' + title, '',
                 '1. 관련: %s 운영 기본계획' % DEPT] + ['%d. %s' % (i + 2, s) for i, s in enumerate(sents)] + [
                 '%d. 협의 담당: %s(%s, %s)' % (len(sents) + 2, who[0], who[1], who[2]),
                 '붙임 세부 일정표 1부. 끝.', '', '한빛군수', '시행 %s (%s)    담당 %s' % (no, date, ME)]
        fname = '%02d_%s_%s.hwpx' % (mon if mon >= 3 else mon + 12, no.replace('-', ''), title.replace(' ', '_'))
        real = write_hwpx(tpl, os.path.join(folder, '01_공문', fname), paras)
        rel = '01_공문/' + real
        made.append(rel)
        add_answer('월별 할 일', todo, '%d년 %d월' % (year, mon), rel, '구역 1', sents[0])
        add_answer('협의할 사람', '%s · %s · %s · %s' % (who[0], who[1], who[2], who[3]), '', rel, '구역 1', '%s(%s, %s)' % (who[0], who[1], who[2]))

    # 2) 엑셀 5
    sched = [['월', '업무', '기한', '관련 문서']] + [['%d월' % m, t, due, no] for (m, no, _, _, due, _, t) in MONTHS]
    write_xlsx(os.path.join(folder, '02_대장·현황(엑셀)', '연간_업무일정표_2025.xlsx'), [('일정', sched)])
    made.append('02_대장·현황(엑셀)/연간_업무일정표_2025.xlsx')
    ledger = [['연번', '일자', '품명', '수량', '금액(원)', '비고']]
    items = ['A4 복사용지', '토너 카트리지', '볼펜(흑)', '결재판', '파일철', '건전지 AA', '포스트잇', '스테이플러 심']
    for i in range(24):
        ledger.append([i + 1, '2025-%02d-%02d' % (3 + i // 2 if 3 + i // 2 <= 12 else (3 + i // 2) - 12, random.randint(1, 28)), random.choice(items), random.randint(1, 20), random.randint(5, 90) * 1000, '영수증 미첨부' if i in (7, 15, 19) else ''])
    write_xlsx(os.path.join(folder, '02_대장·현황(엑셀)', '소모품_구매대장_2025.xlsx'), [('대장', ledger)])
    made.append('02_대장·현황(엑셀)/소모품_구매대장_2025.xlsx')
    contracts = [['계약번호', '계약명', '상대방', '기간', '월 금액(원)', '담당'], ['HB-2023-07', '복합기 임대', '주식회사 가온복합기(가상 업체)', '2023-07-01~2025-06-30', 180000, '최우진 010-0000-0505'], ['HB-2024-02', '정수기 임대', '맑은물렌탈(가상 업체)', '2024-02-01~2026-01-31', 35000, '김예시 010-0000-0909'], ['HB-2024-11', '청사 청소 용역', '깨끗한손(가상 업체)', '2024-11-01~2025-10-31', 1200000, '박모의 010-0000-1010']]
    write_xlsx(os.path.join(folder, '02_대장·현황(엑셀)', '계약_현황.xlsx'), [('계약', contracts)])
    made.append('02_대장·현황(엑셀)/계약_현황.xlsx')
    add_answer('협의할 사람', '최우진 · 주식회사 가온복합기(가상 업체) 영업팀 · 010-0000-0505 · 복합기 임대 계약·소모품', '', '02_대장·현황(엑셀)/계약_현황.xlsx', '시트 「계약」', '최우진 010-0000-0505')
    contacts = [['이름', '소속', '전화', '용건']] + [list(p) for p in PEOPLE]
    write_xlsx(os.path.join(folder, '02_대장·현황(엑셀)', '업무_연락처(가상).xlsx'), [('연락처', contacts)])
    made.append('02_대장·현황(엑셀)/업무_연락처(가상).xlsx')
    budget = [['항목', '예산액(원)', '집행액(원)', '집행률', '기준']] + [['사무관리비', 48000000, 34080000, '71%', '2025-12-31'], ['공공운영비', 36000000, 31320000, '87%', '2025-12-31'], ['여비', 6000000, 5400000, '90%', '2025-12-31']]
    write_xlsx(os.path.join(folder, '02_대장·현황(엑셀)', '예산집행현황_2025.xlsx'), [('집행', budget)])
    made.append('02_대장·현황(엑셀)/예산집행현황_2025.xlsx')
    add_answer('진행 중 현안', ISSUES[3][0] + ': ' + ISSUES[3][1], '2025년 12월', '02_대장·현황(엑셀)/예산집행현황_2025.xlsx', '시트 「집행」', '71%')

    # 3) 워드 4 (+ PDF 3)
    manual = ['서무 업무 매뉴얼(초안)', '1. 소모품 구매', '소모품은 나라장터 종합쇼핑몰에서 구매하고 1회 300만 원을 넘지 않는다.', '구매 즉시 소모품 구매대장에 적고 영수증을 붙인다.', '2. 계약 관리', '임대·용역 계약은 만료 2개월 전에 갱신 여부를 검토한다.', '3. 예산', '분기별 집행 목표 25%를 지키고 매월 말 예산집행현황을 갱신한다.', '4. 기록물', '보존기간이 지난 기록물은 매년 11월에 기록관으로 이관한다.']
    write_docx(os.path.join(folder, '03_계획·매뉴얼(워드)', '서무_업무매뉴얼_초안.docx'), manual)
    made.append('03_계획·매뉴얼(워드)/서무_업무매뉴얼_초안.docx')
    add_answer('월별 할 일', '임대·용역 계약은 만료 2개월 전에 갱신 여부 검토(상시)', '상시', '03_계획·매뉴얼(워드)/서무_업무매뉴얼_초안.docx', '문단 6', '임대·용역 계약은 만료 2개월 전에 갱신 여부를 검토한다.')
    minutes = ['2025년 10월 과 업무회의 회의록', '일시: 2025. 10. 21.(화) 14:00', '장소: 행정지원과 회의실', '참석: 과장, 서무담당 외 4명', '안건 1. 자체감사 지적사항 조치', '소모품 구매대장 누락 3건 중 2건은 영수증을 찾아 보완하였고, 1건은 영수증을 찾지 못해 미조치 상태입니다.', '감사실(윤채원)과 협의하여 11월 중 확인서로 갈음할지 결정하기로 함.', '안건 2. 기록물 이관 준비', '11월 28일 이관 기한을 지키기 위해 10월 말까지 목록 초안을 만들기로 함.']
    write_docx(os.path.join(folder, '03_계획·매뉴얼(워드)', '과_업무회의_회의록_202510.docx'), minutes, table=[['안건', '담당', '기한'], ['감사 지적사항 조치', '서무담당', '2025-11-30'], ['기록물 이관 목록 초안', '서무담당', '2025-10-31']])
    made.append('03_계획·매뉴얼(워드)/과_업무회의_회의록_202510.docx')
    add_answer('진행 중 현안', ISSUES[2][0] + ': ' + ISSUES[2][1], '2025년 10월~', '03_계획·매뉴얼(워드)/과_업무회의_회의록_202510.docx', '문단 6', '1건은 영수증을 찾지 못해 미조치 상태입니다.')
    prev = ['인수인계서(2024년, 전전임자 작성)', '1. 주요 업무: 소모품·계약·예산·기록물·청사 관리 보조', '2. 연간 주요 일정: 3월 소모품 계획, 4월 소방점검, 7월 예산 요구, 9월 행정사무감사, 11월 기록물 이관, 12월 재물조사', '3. 유의사항: 행정사무감사 자료는 제출 1주 전 과장 검토를 받을 것', '4. 비공식 팁: 복합기 업체는 연락이 느리니 전화 뒤 문자로 한 번 더 남길 것']
    write_docx(os.path.join(folder, '03_계획·매뉴얼(워드)', '인수인계서_2024_전전임자.docx'), prev)
    made.append('03_계획·매뉴얼(워드)/인수인계서_2024_전전임자.docx')
    plan = ['2026년 행정지원과 서무 업무 추진계획(안)', '1. 목표: 사무관리비 집행률 95% 이상, 감사 지적 0건', '2. 추진 과제', '가. 소모품 구매대장 월 1회 점검(매월 말)', '나. 계약 만료 2개월 전 알림표 운영', '다. 개인정보 자체점검 6월 실시', '3. 협조 부서: 재정과(예산), 총무과(청사), 정보통신과(보안)']
    write_docx(os.path.join(folder, '03_계획·매뉴얼(워드)', '2026_서무업무_추진계획안.docx'), plan)
    made.append('03_계획·매뉴얼(워드)/2026_서무업무_추진계획안.docx')
    add_answer('월별 할 일', '소모품 구매대장 월 1회 점검(매월 말)', '매월', '03_계획·매뉴얼(워드)/2026_서무업무_추진계획안.docx', '문단 4', '소모품 구매대장 월 1회 점검(매월 말)')
    pdf_dir = os.path.join(folder, '06_기타')
    for name in ['서무_업무매뉴얼_초안.docx', '인수인계서_2024_전전임자.docx', '2026_서무업무_추진계획안.docx']:
        p = docx_to_pdf(os.path.join(folder, '03_계획·매뉴얼(워드)', name), pdf_dir)
        if p:
            made.append('06_기타/' + os.path.basename(p))
    if os.path.exists(os.path.join(pdf_dir, '인수인계서_2024_전전임자.pdf')):
        add_answer('월별 할 일', '행정사무감사 자료는 제출 1주 전 과장 검토', '2025년 9월', '06_기타/인수인계서_2024_전전임자.pdf', '1쪽', '제출 1주 전 과장 검토를 받을 것')

    # 4) 메일 4
    d0 = datetime.datetime(2025, 11, 20, 10, 15, tzinfo=datetime.timezone(datetime.timedelta(hours=9)))
    write_eml(os.path.join(folder, '04_메일', '20251120_기록물_이관목록_반려.eml'), '한빛군청 기록관 <records@example.kr>', '오지영 <seomu@example.kr>', d0, '[반려] 2025년 기록물 이관 목록 서식 오류', '행정지원과 서무담당님,\n\n제출하신 기록물 이관 목록이 서식 오류(보존기간 열 누락)로 반려되었습니다. 재작성하여 11월 28일까지 다시 제출해 주시기 바랍니다.\n문의: 기록관 담당 신가상 063-000-1111\n\n한빛군청 기록관')
    made.append('04_메일/20251120_기록물_이관목록_반려.eml')
    add_answer('진행 중 현안', ISSUES[1][0] + ': ' + ISSUES[1][1], '2025년 11월', '04_메일/20251120_기록물_이관목록_반려.eml', '본문', '서식 오류(보존기간 열 누락)로 반려되었습니다')
    add_answer('협의할 사람', '신가상 · 한빛군청 기록관 · 063-000-1111 · 기록물 이관 목록', '', '04_메일/20251120_기록물_이관목록_반려.eml', '본문', '신가상 063-000-1111')
    d1 = datetime.datetime(2025, 6, 3, 9, 40, tzinfo=datetime.timezone(datetime.timedelta(hours=9)))
    write_eml(os.path.join(folder, '04_메일', '20250603_개인정보_자체점검표_배포.eml'), '한서윤 <security@example.kr>', '각 부서 서무 <all@example.kr>', d1, '2025년 개인정보 보호 자체점검표 배포', '점검표를 첨부합니다. 6월 27일까지 회신 바랍니다.\n정보통신과 보안담당 한서윤 063-000-0606', attach=('자체점검표.xlsx', b'PK\x03\x04placeholder'))
    made.append('04_메일/20250603_개인정보_자체점검표_배포.eml')
    d2 = datetime.datetime(2026, 1, 14, 16, 5, tzinfo=datetime.timezone(datetime.timedelta(hours=9)))
    write_eml(os.path.join(folder, '04_메일', '20260114_복합기_임대료_협의.eml'), '최우진 <sales@example.kr>', '오지영 <seomu@example.kr>', d2, 'RE: 복합기 임대계약 갱신 조건', '안녕하세요. 저희 쪽 제안은 월 임대료 195,000원입니다. 기존 180,000원 유지는 어렵습니다. 검토 후 연락 부탁드립니다.\n가온복합기 영업팀 최우진 010-0000-0505')
    made.append('04_메일/20260114_복합기_임대료_협의.eml')
    add_answer('진행 중 현안', ISSUES[0][0] + ': ' + ISSUES[0][1], '2026년 1월', '04_메일/20260114_복합기_임대료_협의.eml', '본문', '월 임대료 195,000원입니다')
    d3 = datetime.datetime(2025, 9, 1, 11, 0, tzinfo=datetime.timezone(datetime.timedelta(hours=9)))
    write_eml(os.path.join(folder, '04_메일', '20250901_행감자료_요구.eml'), '정하은 <council@example.kr>', '오지영 <seomu@example.kr>', d3, '2025년 행정사무감사 요구자료 안내', '요구자료 목록을 보냅니다. 9월 12일까지 제출해 주세요.\n군의회 사무과 정하은 063-000-0404')
    made.append('04_메일/20250901_행감자료_요구.eml')

    # 5) 메모 4
    memo1 = '[메모] 복합기 계약\n- 2025.6.30 만료 → 임시로 6개월 연장(조건 동일)해 둠. 2026.1월 재협의 중.\n- 업체는 195,000원 요구. 재정과 강민재 담당은 "증액은 3% 이내"라고 함.\n- 결론 못 냄. 후임자가 2월 안에 결정해야 함.'
    open(os.path.join(folder, '05_메모', '메모_복합기계약.txt'), 'w', encoding='utf-8').write(memo1)
    made.append('05_메모/메모_복합기계약.txt')
    add_answer('진행 중 현안', '복합기 임대계약 2026년 2월 안에 갱신 결정 필요(업체 195,000원 요구, 증액 3% 이내 원칙)', '2026년 2월', '05_메모/메모_복합기계약.txt', '전체', '후임자가 2월 안에 결정해야 함')
    memo2 = '[메모] 청사 2층 복도 누수\n2026.1.8 처음 발견. 1.9 총무과 이수현 담당에게 보수 요청(063-000-0202).\n1.22 또 물 고임. 아직 보수 안 됨. 비 오면 확인할 것.'
    open(os.path.join(folder, '05_메모', '메모_누수.txt'), 'w', encoding='cp949').write(memo2)
    made.append('05_메모/메모_누수.txt')
    add_answer('진행 중 현안', ISSUES[4][0] + ': ' + ISSUES[4][1], '2026년 1월~', '05_메모/메모_누수.txt', '전체', '아직 보수 안 됨')
    memo3 = '# 연간 할 일 체크\n- 3월 소모품 계획 / 4월 소방점검 / 6월 개인정보 점검 / 7월 예산요구 / 9월 행감 / 10월 감사조치 / 11월 기록물 / 12월 재물조사 / 2월 집행계획\n- 소방서 박도윤 주임은 오전에 전화 받음(063-000-0303)'
    open(os.path.join(folder, '05_메모', '연간_할일_체크.md'), 'w', encoding='utf-8').write(memo3)
    made.append('05_메모/연간_할일_체크.md')
    memo4 = '비밀번호 메모(가상): 복합기 관리자 0000, 창고 열쇠는 서무 책상 둘째 서랍'
    open(os.path.join(folder, '05_메모', '메모_기타.txt'), 'w', encoding='utf-8').write(memo4)
    made.append('05_메모/메모_기타.txt')

    # 6) 중첩 압축 1, 못 읽을 파일 3
    with zipfile.ZipFile(os.path.join(folder, '06_기타', '행감_제출자료_2025.zip'), 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('제출목록.txt', '1. 사무관리비 집행내역\n2. 계약 현황\n3. 소모품 구매대장\n제출일 2025-09-11, 담당 오지영')
        bio = io.BytesIO()
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = '집행내역'
        ws.append(['월', '집행액(원)']); [ws.append(['%d월' % m, random.randint(2, 5) * 1000000]) for m in range(1, 9)]
        wb.save(bio)
        z.writestr('사무관리비_집행내역.xlsx', bio.getvalue())
    made.append('06_기타/행감_제출자료_2025.zip')
    add_answer('월별 할 일', '행정사무감사 자료 3종(집행내역·계약 현황·구매대장) 9월 제출', '2025년 9월', '06_기타/행감_제출자료_2025.zip ▸ 제출목록.txt', '전체', '제출일 2025-09-11')
    open(os.path.join(folder, '06_기타', '옛공문_스캔.jpg'), 'wb').write(b'\xff\xd8\xff\xe0' + os.urandom(2048))
    made.append('06_기타/옛공문_스캔.jpg')
    open(os.path.join(folder, '06_기타', '손상된_문서.hwp'), 'wb').write(os.urandom(4096))
    made.append('06_기타/손상된_문서.hwp')
    try:
        from pypdf import PdfWriter
        w = PdfWriter(); w.add_blank_page(width=595, height=842); w.encrypt('secret')
        with open(os.path.join(folder, '06_기타', '암호걸린_안내문.pdf'), 'wb') as f:
            w.write(f)
        made.append('06_기타/암호걸린_안내문.pdf')
    except Exception:
        pass

    # 정답표
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = '정답'
    ws.append(['연번', '유형', '내용', '시기', '출처 파일', '출처 위치', '발췌(원문 그대로)'])
    for i, a in enumerate(ANSWERS, 1):
        ws.append([i] + a)
    ws2 = wb.create_sheet('파일목록'); ws2.append(['파일', '읽기 가능'])
    for m in made:
        ws2.append([m, '아니요(의도된 불량 파일)' if m.split('/')[-1] in ('옛공문_스캔.jpg', '손상된_문서.hwp', '암호걸린_안내문.pdf') else '예'])
    wb.save(os.path.join(out, '정답표.xlsx'))
    with open(os.path.join(out, 'README.md'), 'w', encoding='utf-8') as f:
        f.write('# 가상기관 모의 업무폴더(돌다리 시험용)\n\n가상 기관 「%s %s」 서무담당의 1년치(2025.3.~2026.2.) 자료를 합성한 것이다. 기관·사람·전화·업체·문서번호가 모두 가상이며 실제 자료는 없다. 파일 %d개(의도된 불량 파일 3개 포함), 정답 항목 %d개(월별 할 일·협의할 사람·진행 중 현안). 정답표의 「발췌」는 원문과 글자 단위로 같다.\n' % (ORG, DEPT, len(made), len(ANSWERS)))
    print('파일 %d개, 정답 %d개 →' % (len(made), len(ANSWERS)), out)


if __name__ == '__main__':
    main()
